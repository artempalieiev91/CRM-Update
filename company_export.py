"""Формування CSV організацій для CRM (1 рядок на компанію)."""

from __future__ import annotations

import io
from pathlib import Path

import pandas as pd

from industry_mapping import normalize_industry_column
from ids_runtime import log_disk_failure, log_disk_write
from organization_lookup import (
    CRM_LINKEDIN_ID_COLUMN,
    CRM_WEBSITE_ID_COLUMN,
    LINKEDIN_ID_COLUMN,
    OrganizationLookup,
    WEBSITE_ID_COLUMN,
    apply_organization_lookup,
    crm_cell_contains_linkedin,
    crm_cell_contains_website,
    normalize_website_key,
)

ORG_LOOKUP_ISSUE_COLUMN = "Organization lookup issue"
ORG_HIGHLIGHT_COLUMN = "Підсвічити (CRM Website / CRM LinkedIn)"
from excel_utils import excel_cell_value
from state_abbreviations import normalize_state_columns

COMPANY_COLUMN = "Company"
WEBSITE_COLUMN = "Website"
LINKEDIN_COMPANY_COLUMN = "Linkedin Company"
APOLLO_ACCOUNT_ID_COLUMN = "Organization - Apollo Account id"
PIPEDRIVE_ORG_ID_COLUMN = "Organization - ID"

COMPANIES_REVIEW_TAIL_COLUMNS: tuple[str, ...] = (
    ORG_HIGHLIGHT_COLUMN,
    ORG_LOOKUP_ISSUE_COLUMN,
)

COMPANY_CRM_COLUMNS: tuple[str, ...] = (
    "Company",
    "Website",
    "Industry",
    "Country",
    "State",
    "City",
    LINKEDIN_COMPANY_COLUMN,
    "Number of employees",
    PIPEDRIVE_ORG_ID_COLUMN,
    WEBSITE_ID_COLUMN,
    LINKEDIN_ID_COLUMN,
    CRM_WEBSITE_ID_COLUMN,
    CRM_LINKEDIN_ID_COLUMN,
    APOLLO_ACCOUNT_ID_COLUMN,
)

DUPLICATE_NAME_COLUMN = "Duplicate company name"

_MERGE_PROTECTED_COLUMNS = frozenset(
    {
        PIPEDRIVE_ORG_ID_COLUMN,
        WEBSITE_ID_COLUMN,
        LINKEDIN_ID_COLUMN,
        CRM_WEBSITE_ID_COLUMN,
        CRM_LINKEDIN_ID_COLUMN,
        APOLLO_ACCOUNT_ID_COLUMN,
        DUPLICATE_NAME_COLUMN,
        ORG_LOOKUP_ISSUE_COLUMN,
        ORG_HIGHLIGHT_COLUMN,
    }
)
DUPLICATES_MERGE_COLUMNS: tuple[str, ...] = (COMPANY_COLUMN,)
_MERGE_CHANGE_LOG_COLUMNS = (
    APOLLO_ACCOUNT_ID_COLUMN,
    "Column",
    "Before",
    "After",
)
_MERGE_CHANGE_DETAIL_COLUMNS = ("Column", "Before", "After")


def _strip_column_names(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(c).strip().lstrip("\ufeff") for c in out.columns]
    return out


def resolve_apollo_account_id_column(df: pd.DataFrame) -> str | None:
    work = _strip_column_names(df)
    if APOLLO_ACCOUNT_ID_COLUMN in work.columns:
        return APOLLO_ACCOUNT_ID_COLUMN
    for col in work.columns:
        key = str(col).casefold()
        if "apollo" in key and "account" in key and "id" in key:
            return col
    return None


def normalize_merge_changes_df(df: pd.DataFrame | None) -> pd.DataFrame | None:
    """Повертає журнал змін у канонічному вигляді або None, якщо це не журнал."""
    if df is None or not isinstance(df, pd.DataFrame):
        return None
    work = _strip_column_names(df)
    if work.empty:
        return work

    apollo_col = resolve_apollo_account_id_column(work)
    if apollo_col is None or not set(_MERGE_CHANGE_DETAIL_COLUMNS).issubset(work.columns):
        return None

    if apollo_col != APOLLO_ACCOUNT_ID_COLUMN:
        work = work.rename(columns={apollo_col: APOLLO_ACCOUNT_ID_COLUMN})

    return work.loc[:, list(_MERGE_CHANGE_LOG_COLUMNS)].reset_index(drop=True)


def _company_name_key(value: object) -> str:
    return str(value or "").strip().casefold()


def duplicate_company_name_keys(df: pd.DataFrame) -> frozenset[str]:
    """Назви Company, що зустрічаються у 2+ рядках (без урахування регістру)."""
    names = df[COMPANY_COLUMN].map(_company_name_key)
    names = names.loc[names != ""]
    counts = names.value_counts()
    return frozenset(k for k, n in counts.items() if n > 1)


def annotate_duplicate_company_names(df: pd.DataFrame) -> pd.DataFrame:
    dupes = duplicate_company_name_keys(df)
    out = df.copy()
    out[DUPLICATE_NAME_COLUMN] = out[COMPANY_COLUMN].map(
        lambda v: "yes" if _company_name_key(v) in dupes else ""
    )
    return out


def build_duplicate_name_report(df: pd.DataFrame) -> pd.DataFrame:
    """Зведення: одна назва → скільки різних організацій (website / Apollo id)."""
    dupes = duplicate_company_name_keys(df)
    if not dupes:
        return pd.DataFrame(columns=["Company", "Count", "Websites", "Apollo Account ids"])

    rows: list[dict[str, object]] = []
    for name_display, group in df.groupby(
        df[COMPANY_COLUMN].map(lambda v: _company_name_key(v)), sort=False
    ):
        if name_display not in dupes:
            continue
        company_label = str(group[COMPANY_COLUMN].iloc[0]).strip()
        websites = sorted({str(w).strip() for w in group[WEBSITE_COLUMN] if str(w).strip()})
        apollo_ids = sorted(
            {str(a).strip() for a in group[APOLLO_ACCOUNT_ID_COLUMN] if str(a).strip()}
        )
        rows.append(
            {
                "Company": company_label,
                "Count": len(group),
                "Websites": "; ".join(websites),
                "Apollo Account ids": "; ".join(apollo_ids),
            }
        )
    report = pd.DataFrame(rows)
    return report.sort_values(["Count", "Company"], ascending=[False, True]).reset_index(drop=True)

_DEDUP_KEY_COLUMNS = (APOLLO_ACCOUNT_ID_COLUMN, WEBSITE_COLUMN, COMPANY_COLUMN)


def _fill_na_str(df: pd.DataFrame) -> pd.DataFrame:
    return df.fillna("").astype(str).replace({"nan": "", "<NA>": ""})


def _company_dedup_key(row: pd.Series) -> str:
    """
    Ключ дедуплікації з дослідження: спочатку website (1 домен = 1 компанія),
    інакше Apollo Account id, інакше назва Company.
    """
    website = normalize_website_key(row.get(WEBSITE_COLUMN, ""))
    if website:
        return f"website:{website}"
    apollo = str(row.get(APOLLO_ACCOUNT_ID_COLUMN, "")).strip()
    if apollo:
        return f"apollo:{apollo}"
    company = str(row.get(COMPANY_COLUMN, "")).strip().casefold()
    if company:
        return f"company:{company}"
    return ""


def dedupe_one_row_per_company(df: pd.DataFrame) -> pd.DataFrame:
    """Залишає перший рядок на компанію (website → Apollo Account id → назва)."""
    src = _fill_na_str(df)
    keys: list[str] = []
    for _, row in src.iterrows():
        keys.append(_company_dedup_key(row))
    work = src.copy()
    work["_dedup_key"] = keys
    work = work.loc[work["_dedup_key"] != ""]
    return work.drop_duplicates(subset=["_dedup_key"], keep="first").drop(columns=["_dedup_key"])


def align_companies_to_crm_schema(df: pd.DataFrame) -> pd.DataFrame:
    """Додає відсутні колонки Comp-шаблону (напр. після старого збереження workspace)."""
    out = df.copy()
    for col in COMPANY_CRM_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    ordered = list(COMPANY_CRM_COLUMNS)
    if DUPLICATE_NAME_COLUMN in out.columns and DUPLICATE_NAME_COLUMN not in ordered:
        ordered.append(DUPLICATE_NAME_COLUMN)
    for col in COMPANIES_REVIEW_TAIL_COLUMNS:
        if col in out.columns and col not in ordered:
            ordered.append(col)
    tail = [c for c in out.columns if c not in ordered]
    return out[ordered + tail]


def build_company_export(
    df: pd.DataFrame,
    *,
    org_lookup: OrganizationLookup | None = None,
) -> pd.DataFrame:
    unique = dedupe_one_row_per_company(df)
    out = pd.DataFrame()
    for col in COMPANY_CRM_COLUMNS:
        if col in (
            PIPEDRIVE_ORG_ID_COLUMN,
            WEBSITE_ID_COLUMN,
            LINKEDIN_ID_COLUMN,
            CRM_WEBSITE_ID_COLUMN,
            CRM_LINKEDIN_ID_COLUMN,
        ):
            if col in unique.columns:
                out[col] = unique[col]
            else:
                out[col] = ""
        else:
            out[col] = unique[col] if col in unique.columns else ""

    out = normalize_state_columns(out, columns=("State",))
    out = normalize_industry_column(out)
    out = apply_organization_lookup(out, org_lookup)
    out = out[list(COMPANY_CRM_COLUMNS)]
    out = annotate_duplicate_company_names(out)
    out["_sort_name"] = out[COMPANY_COLUMN].map(_company_name_key)
    out = out.sort_values(
        ["_sort_name", WEBSITE_COLUMN, APOLLO_ACCOUNT_ID_COLUMN],
        kind="stable",
    ).drop(columns=["_sort_name"])
    return out.reset_index(drop=True)


def save_company_export(
    df: pd.DataFrame,
    path: Path | str,
    *,
    org_lookup: OrganizationLookup | None = None,
) -> Path:
    out_path = Path(path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        companies_for_crm_file(build_company_export(df, org_lookup=org_lookup)).to_csv(
            out_path, index=False, encoding="utf-8-sig"
        )
    except OSError as exc:
        log_disk_failure(out_path, exc, action="write companies CSV")
        raise
    log_disk_write(out_path, action="write companies CSV")
    return out_path


def _has_org_contact_value(value: object) -> bool:
    return str(value).strip() != ""


def _crm_org_id(value: object) -> bool:
    return str(value).strip().isdigit()


def _has_crm_url_value(value: object) -> bool:
    return str(value).strip() != ""


def _trimmed_cells_differ(a: object, b: object) -> bool:
    """Як TRIM(L)<>TRIM(B) у Google Sheets (посимвольно, з урахуванням регістру)."""
    return str(a).strip() != str(b).strip()


def _crm_website_highlight_mismatch(website: object, crm_website: object) -> bool:
    """
    Підсвітка L: CRM не порожній і Website не входить у список CRM (через кому).
    """
    if not _has_crm_url_value(crm_website):
        return False
    if not _has_org_contact_value(website):
        return _trimmed_cells_differ(crm_website, website)
    return not crm_cell_contains_website(crm_website, website)


def _crm_linkedin_highlight_mismatch(linkedin: object, crm_linkedin: object) -> bool:
    """CRM Linkedin ID: Linkedin Company не входить у список CRM (через кому)."""
    if not _has_crm_url_value(crm_linkedin):
        return False
    if not _has_org_contact_value(linkedin):
        return _trimmed_cells_differ(crm_linkedin, linkedin)
    return not crm_cell_contains_linkedin(crm_linkedin, linkedin)


def _website_linkedin_id_highlight_mismatch(wid: object, lid: object) -> bool:
    """Website ID і Linkedin ID обидва заповнені, але не збігаються."""
    w = str(wid).strip()
    l = str(lid).strip()
    return bool(w and l and w != l)


def organization_highlight_issue_parts(row: pd.Series) -> list[str]:
    """
    Як умовне форматування Comp (кол. L, M), з нормалізацією URL для доменів/LinkedIn:
    AND(CRM <> \"\", значення реально різні), порожнє CRM не підсвічується.
    """
    issues: list[str] = []
    web = str(row.get(WEBSITE_COLUMN, "")).strip()
    li = str(row.get(LINKEDIN_COMPANY_COLUMN, "")).strip()
    crm_web = str(row.get(CRM_WEBSITE_ID_COLUMN, "")).strip()
    crm_li = str(row.get(CRM_LINKEDIN_ID_COLUMN, "")).strip()

    if _crm_website_highlight_mismatch(web, crm_web):
        issues.append("Підсвічено: CRM Website ID")
    if _crm_linkedin_highlight_mismatch(li, crm_li):
        issues.append("Підсвічено: CRM Linkedin ID")

    wid = str(row.get(WEBSITE_ID_COLUMN, "")).strip()
    lid = str(row.get(LINKEDIN_ID_COLUMN, "")).strip()
    if _website_linkedin_id_highlight_mismatch(wid, lid):
        issues.append("Підсвічено: Website ID")
        issues.append("Підсвічено: Linkedin ID")

    return issues


def organization_extra_issue_parts(row: pd.Series) -> list[str]:
    """Інші сигнали перевірки (без кольорової підсвітки CRM URL)."""
    issues: list[str] = []
    web = str(row.get(WEBSITE_COLUMN, "")).strip()
    li = str(row.get(LINKEDIN_COMPANY_COLUMN, "")).strip()
    oid = str(row.get(PIPEDRIVE_ORG_ID_COLUMN, "")).strip()
    wid = str(row.get(WEBSITE_ID_COLUMN, "")).strip()
    lid = str(row.get(LINKEDIN_ID_COLUMN, "")).strip()

    if web and not _crm_org_id(wid):
        issues.append("Website не знайдено в CRM")
    if li and not _crm_org_id(lid):
        issues.append("LinkedIn не знайдено в CRM")
    if (web or li) and not oid:
        issues.append("немає Organization - ID")
    if oid and wid and lid and oid not in (wid, lid):
        issues.append("Organization - ID не збігається з Website/Linkedin ID")
    crm_web = str(row.get(CRM_WEBSITE_ID_COLUMN, "")).strip()
    crm_li = str(row.get(CRM_LINKEDIN_ID_COLUMN, "")).strip()
    if oid and not crm_web and not crm_li:
        issues.append("є ID, але немає CRM Website/Linkedin у базі")

    return issues


def organization_row_is_highlighted(row: pd.Series) -> bool:
    return bool(organization_highlight_issue_parts(row))


def organization_lookup_issue_label(row: pd.Series) -> str:
    """Причини перевірки: підсвітка CRM URL + інші сигнали формул на Comp."""
    return "; ".join(
        organization_highlight_issue_parts(row) + organization_extra_issue_parts(row)
    )


def _annotate_companies_review(work: pd.DataFrame) -> pd.DataFrame:
    out = work.copy()

    def _annotate_row(row: pd.Series) -> pd.Series:
        highlight = organization_highlight_issue_parts(row)
        issue = "; ".join(highlight + organization_extra_issue_parts(row))
        return pd.Series(
            {
                ORG_LOOKUP_ISSUE_COLUMN: issue,
                ORG_HIGHLIGHT_COLUMN: "yes" if highlight else "",
            }
        )

    labels = out.apply(_annotate_row, axis=1)
    out[ORG_LOOKUP_ISSUE_COLUMN] = labels[ORG_LOOKUP_ISSUE_COLUMN]
    out[ORG_HIGHLIGHT_COLUMN] = labels[ORG_HIGHLIGHT_COLUMN]
    return out


def _fix_zurich_city_for_review(work: pd.DataFrame) -> pd.DataFrame:
    """State = Zurich і City = Zurich → City = Zuerich (companies-review)."""
    out = work.copy()
    if "State" not in out.columns or "City" not in out.columns:
        return out
    state = out["State"].astype(str).str.strip()
    city = out["City"].astype(str).str.strip()
    mask = state.str.casefold().eq("zurich") & city.str.casefold().eq("zurich")
    if mask.any():
        out.loc[mask, "City"] = "Zuerich"
    return out


def prepare_companies_review(df: pd.DataFrame) -> pd.DataFrame:
    """Повна таблиця Comp + колонки перевірки (для кешу та Excel)."""
    work = align_companies_to_crm_schema(df)
    work = _fix_zurich_city_for_review(work)
    work = _annotate_companies_review(work)
    ordered = list(COMPANY_CRM_COLUMNS)
    if DUPLICATE_NAME_COLUMN in work.columns and DUPLICATE_NAME_COLUMN not in ordered:
        ordered.append(DUPLICATE_NAME_COLUMN)
    ordered.extend(COMPANIES_REVIEW_TAIL_COLUMNS)
    tail = [c for c in work.columns if c not in ordered]
    return work[ordered + tail].reset_index(drop=True)


def companies_review_export(df: pd.DataFrame) -> pd.DataFrame:
    """Усі компанії в порядку Comp + «Підсвічити» та «Organization lookup issue»."""
    return prepare_companies_review(df)


def organization_review_stats(review_df: pd.DataFrame) -> dict[str, int]:
    """Метрики Organization ID з уже підготовленої review-таблиці."""
    total = len(review_df)
    if total == 0:
        return {
            "total": 0,
            "with_org_id": 0,
            "with_website_id": 0,
            "with_linkedin_id": 0,
            "with_crm_website": 0,
            "with_crm_linkedin": 0,
            "issues": 0,
            "highlighted": 0,
        }

    def _count(col: str) -> int:
        if col not in review_df.columns:
            return 0
        return int(review_df[col].astype(str).str.strip().ne("").sum())

    issues = int(review_df[ORG_LOOKUP_ISSUE_COLUMN].astype(str).str.strip().ne("").sum()) if ORG_LOOKUP_ISSUE_COLUMN in review_df.columns else 0
    highlighted = int((review_df[ORG_HIGHLIGHT_COLUMN].astype(str).str.strip() == "yes").sum()) if ORG_HIGHLIGHT_COLUMN in review_df.columns else 0
    return {
        "total": total,
        "with_org_id": _count(PIPEDRIVE_ORG_ID_COLUMN),
        "with_website_id": _count(WEBSITE_ID_COLUMN),
        "with_linkedin_id": _count(LINKEDIN_ID_COLUMN),
        "with_crm_website": _count(CRM_WEBSITE_ID_COLUMN),
        "with_crm_linkedin": _count(CRM_LINKEDIN_ID_COLUMN),
        "issues": issues,
        "highlighted": highlighted,
    }


def organization_highlight_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Рядки з підсвіткою CRM Website ID / CRM Linkedin ID (як у Google Sheets)."""
    full = companies_review_export(df)
    return full.loc[full[ORG_HIGHLIGHT_COLUMN] == "yes"].copy().reset_index(drop=True)


def _organization_cell_highlight_kind(issue: str, column: str) -> str | None:
    if not issue:
        return None
    if column == CRM_WEBSITE_ID_COLUMN and "Підсвічено: CRM Website ID" in issue:
        return "crm"
    if column == CRM_LINKEDIN_ID_COLUMN and "Підсвічено: CRM Linkedin ID" in issue:
        return "crm"
    if column == WEBSITE_ID_COLUMN and "Підсвічено: Website ID" in issue:
        return "id"
    if column == LINKEDIN_ID_COLUMN and "Підсвічено: Linkedin ID" in issue:
        return "id"
    return None


def _organization_cell_highlight_fill(issue: str, column: str) -> bool:
    return _organization_cell_highlight_kind(issue, column) is not None


def companies_crm_file_columns(df: pd.DataFrame) -> list[str]:
    """Колонки файлу для імпорту в CRM (14 колонок Comp, без службових)."""
    return [c for c in COMPANY_CRM_COLUMNS if c in df.columns]


def companies_for_crm_file(df: pd.DataFrame) -> pd.DataFrame:
    """Таблиця companies лише з колонками Comp (без Duplicate company name тощо)."""
    cols = companies_crm_file_columns(df)
    return df.loc[:, cols].reset_index(drop=True)


def companies_review_xlsx_columns(df: pd.DataFrame) -> list[str]:
    """Колонки Excel — як Comp для CRM, без службових полів перевірки."""
    return companies_crm_file_columns(df)


def companies_review_to_xlsx_bytes(df: pd.DataFrame) -> bytes:
    """Excel з підсвіткою CRM URL (червоний) та Website/Linkedin ID (рожевий)."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    crm_fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
    id_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
    wb = Workbook()
    ws = wb.active
    ws.title = "Companies review"

    columns = companies_review_xlsx_columns(df)
    ws.append(columns)
    for _, row in df.iterrows():
        issue = str(row.get(ORG_LOOKUP_ISSUE_COLUMN, "")).strip()
        ws.append([excel_cell_value(c, row.get(c, "")) for c in columns])
        if not issue:
            continue
        excel_row = ws.max_row
        for col_idx, col_name in enumerate(columns, start=1):
            kind = _organization_cell_highlight_kind(issue, col_name)
            if kind == "crm":
                ws.cell(row=excel_row, column=col_idx).fill = crm_fill
            elif kind == "id":
                ws.cell(row=excel_row, column=col_idx).fill = id_fill

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def organization_id_stats(df: pd.DataFrame) -> dict[str, int]:
    """Підсумок по ID-колонках для UI."""
    total = len(df)
    if total == 0:
        return {
            "total": 0,
            "with_org_id": 0,
            "with_website_id": 0,
            "with_linkedin_id": 0,
            "with_crm_website": 0,
            "with_crm_linkedin": 0,
            "issues": 0,
            "highlighted": 0,
        }

    def _count(col: str) -> int:
        if col not in df.columns:
            return 0
        return int(df[col].astype(str).str.strip().ne("").sum())

    if ORG_LOOKUP_ISSUE_COLUMN in df.columns and ORG_HIGHLIGHT_COLUMN in df.columns:
        issues = int(df[ORG_LOOKUP_ISSUE_COLUMN].astype(str).str.strip().ne("").sum())
        highlighted = int((df[ORG_HIGHLIGHT_COLUMN].astype(str).str.strip() == "yes").sum())
    else:
        review = prepare_companies_review(df)
        issues = int(review[ORG_LOOKUP_ISSUE_COLUMN].astype(str).str.strip().ne("").sum())
        highlighted = int((review[ORG_HIGHLIGHT_COLUMN].astype(str).str.strip() == "yes").sum())
    return {
        "total": total,
        "with_org_id": _count(PIPEDRIVE_ORG_ID_COLUMN),
        "with_website_id": _count(WEBSITE_ID_COLUMN),
        "with_linkedin_id": _count(LINKEDIN_ID_COLUMN),
        "with_crm_website": _count(CRM_WEBSITE_ID_COLUMN),
        "with_crm_linkedin": _count(CRM_LINKEDIN_ID_COLUMN),
        "issues": issues,
        "highlighted": highlighted,
    }


def organization_problem_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Рядки companies з проблемами метчу організацій (для вигрузки та правок)."""
    work = df.copy()
    work[ORG_LOOKUP_ISSUE_COLUMN] = work.apply(organization_lookup_issue_label, axis=1)
    out = work.loc[work[ORG_LOOKUP_ISSUE_COLUMN].astype(str).str.strip() != ""].copy()
    return out.sort_values(
        [COMPANY_COLUMN, WEBSITE_COLUMN, APOLLO_ACCOUNT_ID_COLUMN],
        kind="stable",
    ).reset_index(drop=True)


def duplicate_company_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Рядки, де одна й та сама назва Company (без урахування регістру) у 2+ рядках."""
    work = annotate_duplicate_company_names(df)
    mask = work[DUPLICATE_NAME_COLUMN].astype(str).str.strip().str.casefold() == "yes"
    out = work.loc[mask].copy()
    return out.sort_values(
        [COMPANY_COLUMN, WEBSITE_COLUMN, APOLLO_ACCOUNT_ID_COLUMN],
        kind="stable",
    ).reset_index(drop=True)


def format_merge_changes_summary(changes_df: pd.DataFrame) -> pd.DataFrame:
    """Коротке превʼю: один рядок на Apollo id."""
    normalized = normalize_merge_changes_df(changes_df)
    if normalized is None or normalized.empty:
        return pd.DataFrame(columns=[APOLLO_ACCOUNT_ID_COLUMN, "Зміни"])

    rows: list[dict[str, str]] = []
    for apollo, group in normalized.groupby(APOLLO_ACCOUNT_ID_COLUMN, sort=False):
        parts = [
            f"{str(r['Column']).strip()}: «{r['Before']}» → «{r['After']}»"
            for _, r in group.iterrows()
        ]
        rows.append({APOLLO_ACCOUNT_ID_COLUMN: apollo, "Зміни": "; ".join(parts)})
    return pd.DataFrame(rows)


def merge_companies_by_apollo_id(
    base_df: pd.DataFrame,
    edited_df: pd.DataFrame,
    *,
    merge_columns: tuple[str, ...] | None = None,
) -> tuple[pd.DataFrame, dict[str, int], pd.DataFrame]:
    """Оновлює companies за Organization - Apollo Account id з відредагованого CSV.

    ID-колонки (Organization - ID, Website ID, …) ніколи не перезаписуються з CSV.
    Якщо передано merge_columns — застосовуються лише ці поля (напр. Company для дублікатів).
    """
    if APOLLO_ACCOUNT_ID_COLUMN not in edited_df.columns:
        raise ValueError(
            f"У завантаженому файлі немає колонки «{APOLLO_ACCOUNT_ID_COLUMN}»."
        )
    if merge_columns is not None:
        missing = [c for c in merge_columns if c not in edited_df.columns]
        if missing:
            raise ValueError(
                f"У завантаженому файлі немає колонок для злиття: {', '.join(missing)}."
            )

    base = _fill_na_str(base_df)
    edited = _fill_na_str(edited_df)

    index_by_apollo: dict[str, int] = {}
    for idx in base.index:
        apollo = str(base.at[idx, APOLLO_ACCOUNT_ID_COLUMN]).strip()
        if apollo:
            index_by_apollo[apollo] = idx

    out = base.copy()
    updated = 0
    skipped_no_id = 0
    not_found = 0
    change_rows: list[dict[str, str]] = []

    for _, row in edited.iterrows():
        apollo = str(row[APOLLO_ACCOUNT_ID_COLUMN]).strip()
        if not apollo:
            skipped_no_id += 1
            continue
        if apollo not in index_by_apollo:
            not_found += 1
            continue
        idx = index_by_apollo[apollo]
        if merge_columns is not None:
            cols_to_apply = merge_columns
        else:
            cols_to_apply = tuple(
                c
                for c in edited.columns
                if c not in _MERGE_PROTECTED_COLUMNS and c in out.columns
            )
        for col in cols_to_apply:
            if col not in out.columns:
                continue
            before = str(base.at[idx, col]).strip()
            after = str(row[col]).strip()
            if before != after:
                change_rows.append(
                    {
                        APOLLO_ACCOUNT_ID_COLUMN: apollo,
                        "Column": col,
                        "Before": before,
                        "After": after,
                    }
                )
            out.at[idx, col] = row[col]
        updated += 1

    out = annotate_duplicate_company_names(out)
    out["_sort_name"] = out[COMPANY_COLUMN].map(_company_name_key)
    out = out.sort_values(
        ["_sort_name", WEBSITE_COLUMN, APOLLO_ACCOUNT_ID_COLUMN],
        kind="stable",
    ).drop(columns=["_sort_name"])
    out = out.reset_index(drop=True)

    changes_df = pd.DataFrame(change_rows, columns=list(_MERGE_CHANGE_LOG_COLUMNS))
    changed_apollo = (
        frozenset(changes_df[APOLLO_ACCOUNT_ID_COLUMN].unique())
        if not changes_df.empty
        else frozenset()
    )
    stats = {
        "edited_rows": len(edited),
        "updated": updated,
        "not_found": not_found,
        "skipped_no_id": skipped_no_id,
        "changed_fields": len(changes_df),
        "changed_rows": len(changed_apollo),
    }
    return out, stats, changes_df
